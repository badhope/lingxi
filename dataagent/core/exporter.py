"""
DataAgent - 导出模块
支持 PDF, Excel, CSV 等格式
"""
import json
import csv
from pathlib import Path
from typing import Dict, List, Optional, Any, Union
from datetime import datetime
from core.logging import get_logger

logger = get_logger("exporter")


class Exporter:
    """
    数据导出器
    支持多种格式：CSV, JSON, Excel, PDF
    """
    
    def __init__(self, output_dir: Optional[Path] = None):
        if output_dir is None:
            output_dir = Path.home() / ".dataagent" / "exports"
        
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        # 检查是否有可选库
        self._has_pandas = False
        self._has_excel = False
        self._has_pdf = False
        
        try:
            import pandas
            self._has_pandas = True
        except ImportError:
            pass
        
        try:
            import openpyxl
            self._has_excel = True
        except ImportError:
            pass
        
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
            from reportlab.lib.styles import getSampleStyleSheet
            self._has_pdf = True
        except ImportError:
            pass
        
        logger.info("Exporter initialized")
    
    def get_supported_formats(self) -> List[str]:
        """获取支持的导出格式"""
        formats = ["json", "csv"]
        
        if self._has_excel:
            formats.append("xlsx")
        
        if self._has_pdf:
            formats.append("pdf")
        
        return formats
    
    def export_json(self, data: Union[List[Dict], Dict], filename: Optional[str] = None) -> Path:
        """导出为 JSON"""
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"export_{timestamp}.json"
        
        file_path = self.output_dir / filename
        
        with open(file_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"Exported JSON: {file_path}")
        return file_path
    
    def export_csv(self, data: List[Dict], filename: Optional[str] = None, 
                  headers: Optional[List[str]] = None) -> Path:
        """导出为 CSV"""
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"export_{timestamp}.csv"
        
        file_path = self.output_dir / filename
        
        if not data:
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write("")
            return file_path
        
        # 获取表头
        if headers is None:
            headers = list(data[0].keys())
        
        with open(file_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            for row in data:
                writer.writerow({k: str(row.get(k, '')) for k in headers})
        
        logger.info(f"Exported CSV: {file_path}")
        return file_path
    
    def export_excel(self, data: List[Dict], filename: Optional[str] = None, 
                    sheet_name: str = "Data") -> Optional[Path]:
        """导出为 Excel"""
        if not self._has_excel:
            logger.warning("Excel export not available - install openpyxl")
            return None
        
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
            
            if filename is None:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"export_{timestamp}.xlsx"
            
            file_path = self.output_dir / filename
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
            
            if data:
                headers = list(data[0].keys())
                
                # 写表头
                for col, header in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=col, value=header)
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center")
                
                # 写数据
                for row_idx, row_data in enumerate(data, start=2):
                    for col_idx, header in enumerate(headers, start=1):
                        value = row_data.get(header, '')
                        ws.cell(row=row_idx, column=col_idx, value=str(value))
                
                # 自动调整列宽
                for col in ws.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column].width = adjusted_width
            
            wb.save(file_path)
            logger.info(f"Exported Excel: {file_path}")
            return file_path
            
        except Exception as e:
            logger.error(f"Excel export failed: {e}")
            return None
    
    def export_pdf(self, data: List[Dict], filename: Optional[str] = None,
                  title: str = "Data Report") -> Optional[Path]:
        """导出为 PDF"""
        if not self._has_pdf:
            logger.warning("PDF export not available - install reportlab")
            return None
        
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            
            if filename is None:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"export_{timestamp}.pdf"
            
            file_path = self.output_dir / filename
            
            doc = SimpleDocTemplate(str(file_path), pagesize=letter)
            elements = []
            
            styles = getSampleStyleSheet()
            
            # 标题
            title_para = Paragraph(title, styles['Title'])
            elements.append(title_para)
            elements.append(Spacer(1, 20))
            
            if data:
                headers = list(data[0].keys())
                table_data = [headers]
                
                for row in data:
                    table_data.append([str(row.get(header, '')) for header in headers])
                
                table = Table(table_data)
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ]))
                
                elements.append(table)
            
            doc.build(elements)
            logger.info(f"Exported PDF: {file_path}")
            return file_path
            
        except Exception as e:
            logger.error(f"PDF export failed: {e}")
            return None
    
    def export(self, data: Union[List[Dict], Dict], format: str = "json", **kwargs) -> Optional[Path]:
        """通用导出函数"""
        format = format.lower()
        
        if format == "json":
            return self.export_json(data, kwargs.get("filename"))
        elif format == "csv":
            if isinstance(data, dict):
                data = [data]
            return self.export_csv(data, kwargs.get("filename"), kwargs.get("headers"))
        elif format == "xlsx":
            if isinstance(data, dict):
                data = [data]
            return self.export_excel(data, kwargs.get("filename"), kwargs.get("sheet_name", "Data"))
        elif format == "pdf":
            if isinstance(data, dict):
                data = [data]
            return self.export_pdf(data, kwargs.get("filename"), kwargs.get("title", "Data Report"))
        else:
            logger.error(f"Unsupported format: {format}")
            return None


# 全局导出器
_exporter: Optional[Exporter] = None


def get_exporter() -> Exporter:
    """获取导出器"""
    global _exporter
    if _exporter is None:
        _exporter = Exporter()
    return _exporter
